"""
labeled_docx_comments.py

Build a Word (.docx) document that keeps the pipeline's RESULTS chunk format and,
for every extracted reaction, HIGHLIGHTS each sentence it was pulled from and
attaches a Word COMMENT to each one naming the reaction.

The location hints come from the reaction's `evidence` field in the extraction
JSON: each evidence string is a sentence copied from the chunk text, so we search
for it inside the chunks, highlight it, and anchor its own comment to it. A
reaction cited by three sentences therefore produces three comments, labelled
"evidence 1/3" and so on — a highlight with the wrong reaction on it is then
visible where it sits, instead of being one unexplained highlight among many.
Only the evidence sentences are highlighted; text between them is left plain.

Only MERGED JSON is annotated. A raw extraction file still carries the per-chunk
duplicates the merge collapses, so a report built from one shows the same reaction
several times under several names — which is exactly what a curator should not
have to reconcile by hand. With no argument the newest results/*_merged.json is
used, so the report always reflects the most recent merge rather than whichever
path happened to be typed.

Nothing is pasted in by hand. The merged JSON records the paper it came from
(`source`), so the Results text is re-fetched with the pipeline's own fetcher and
re-split with run_extraction's own splitter OBJECT — the chunks below are the
chunks the model actually saw, not a re-typed approximation of them.

Usage:
    python labeled_docx_comments.py                        # newest *_merged.json
    python labeled_docx_comments.py --source PINK1.pdf     # newest merge OF THAT PAPER
    python labeled_docx_comments.py results/pink1_2prev1next_merged.json
    python labeled_docx_comments.py --list                 # what is available, newest first
    python labeled_docx_comments.py <merged.json> --text /tmp/results.txt
    python labeled_docx_comments.py <merged.json> --refresh -v
Output:
    results/<merged json stem>_annotated.docx
    (derived from the input, so one paper's merge can never land on another's report)
"""

import argparse
import json
import os
import re
import sys
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

ROOT = Path(__file__).resolve().parent
RESULTS_DIR = ROOT / "results"
# Results text is cached per source so a re-run does not repeat the PDF section
# detection (an LLM call chain) or re-hit PubMed Central. --refresh busts it.
CACHE_DIR = RESULTS_DIR / "cache"


def normalize(text: str) -> str:
    """Rejoin hyphenated line-wraps, drop newlines, collapse whitespace.

    Also drops the control characters PDF extraction leaves behind (form feeds at
    page breaks, mostly): python-docx writes runs straight into XML, which permits
    no C0/C1 control except tab/LF/CR. Both the chunk text and the evidence pass
    through here, so offsets stay aligned.
    """
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", " ", text)
    text = re.sub(r"-\n", "", text)        # word-\nwrap -> wordwrap
    text = text.replace("\n", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _merge(blocks, gap=3):
    """Merge (start, end) ranges separated by at most `gap` characters.

    Repairing a PDF word-wrap leaves 0–1 character seams inside one sentence;
    merging those keeps a highlight continuous, while a real interruption
    (a figure citation dropped mid-sentence) stays a separate block and so
    stays unhighlighted.
    """
    out = []
    for s, e in sorted(blocks):
        if out and s - out[-1][1] <= gap:
            out[-1][1] = max(out[-1][1], e)
        else:
            out.append([s, e])
    return [(s, e) for s, e in out]


def find_span(evidence: str, hay: str):
    """
    Locate `evidence` inside `hay` (both normalized).
    Returns (start, end, quality, blocks) or None.
    quality: 1.0 exact, else fraction of evidence chars aligned (fuzzy).
    blocks: the sub-ranges to HIGHLIGHT. For a fuzzy match these are the pieces
    that actually aligned, so text sitting between them — which belongs to no
    evidence sentence — is left unhighlighted even though the comment anchors
    across the whole span.
    """
    ev = normalize(evidence)
    if not ev:
        return None

    idx = hay.find(ev)
    if idx != -1:
        return idx, idx + len(ev), 1.0, [(idx, idx + len(ev))]

    # Fuzzy fallback for PDF word-wraps that lost their hyphen
    # (e.g. "phosphor\nylation" -> "phosphor ylation" != "phosphorylation").
    sm = SequenceMatcher(None, hay, ev, autojunk=False)
    blocks = [b for b in sm.get_matching_blocks() if b.size >= 4]
    if not blocks:
        return None
    matched = sum(b.size for b in blocks)
    if matched < 0.6 * len(ev):
        return None
    start = blocks[0].a
    end = blocks[-1].a + blocks[-1].size
    # Reject a span far longer than the excerpt itself: that means the matching
    # fragments are scattered through the chunk rather than forming the sentence,
    # so even the comment anchor would cover unrelated prose. Word-wrap repair only
    # inflates a true match slightly, so 1.5x leaves ample slack.
    if end - start > 1.5 * len(ev):
        return None
    return (start, end, matched / len(ev),
            _merge((b.a, b.a + b.size) for b in blocks))


def build_comment(rxn: dict, ev_num=None, ev_total=None) -> str:
    """Human-readable comment body describing the reaction.

    One comment is attached per evidence sentence, so a reaction cited by three
    sentences produces three of these. ev_num/ev_total label which sentence this
    is, so a curator reading one margin note knows the reaction rests on others
    too. Summation is deliberately omitted: it is the longest field and repeating
    a paragraph across every evidence sentence is what makes the margin unreadable.
    """
    ca = rxn.get("catalystActivity") or {}
    if isinstance(ca, list):                        # merged records can carry several
        ca = ca[0] if ca and isinstance(ca[0], dict) else {}
    tag = f" (evidence {ev_num}/{ev_total})" if ev_total and ev_total > 1 else ""
    lines = [
        f"REACTION{tag}: {rxn.get('name', '(unnamed)')}",
        f"Type: {rxn.get('reactionType', '?')}"
        f"  |  Confidence: {rxn.get('confidence', '?')}",
        f"Input: {', '.join(rxn.get('input') or []) or '—'}",
        f"Output: {', '.join(rxn.get('output') or []) or '—'}",
    ]
    if ca.get("catalyst"):
        mf = ca.get("molecularFunction") or ""
        lines.append(f"Catalyst: {ca['catalyst']}" + (f" ({mf})" if mf else ""))
    reg = rxn.get("regulatedBy") or []
    if reg:
        reg_str = "; ".join(
            f"{r.get('regulator')} [{r.get('regulationType')}]" for r in reg
        )
        lines.append(f"Regulated by: {reg_str}")
    # merged JSON records which per-chunk reactions were folded together
    merged = rxn.get("merged_names") or []
    if merged:
        lines.append(f"Merged from: {'; '.join(merged)}")
    # context_used is now a list (one reaction can consult several directions);
    # older extractions wrote a bare string, so accept both
    ctx = rxn.get("context_used")
    if isinstance(ctx, list):
        ctx = ", ".join(c for c in ctx if c and c != "none")
    if ctx and ctx != "none":
        lines.append(f"(extraction used context: {ctx})")
    return "\n".join(lines)


def attach_comment(doc, runs, text, author="reaction-extractor", initials="RX"):
    """Attach a comment to the given runs, or fall back to appending comment text."""
    if hasattr(doc, "add_comment"):
        try:
            doc.add_comment(runs, text=text, author=author, initials=initials)
            return
        except Exception:
            pass

    if not runs:
        return
    fallback = text.replace("\n", "  |  ")
    runs[-1].add_text(f" [{author}: {fallback}]")


def locate_evidence(evidence: str, norm_chunks, prefer_chunk=None):
    """
    Find `evidence` across ALL chunks; return (chunk_idx, start, end, quality, blocks)
    for the best match, or None.

    The whole document is searched rather than only the record's `chunk_index`:
    a reaction may be reported in one chunk but evidenced in the previous/next
    chunk the model was shown as context, and merged records carry no chunk index
    at all. `prefer_chunk` breaks ties in favour of the recorded chunk.
    """
    best = None  # (quality, is_preferred, chunk_idx, start, end, blocks)
    for ci, hay in enumerate(norm_chunks):
        span = find_span(evidence, hay)
        if span is None:
            continue
        s, e, q, blocks = span
        key = (q, ci == prefer_chunk)
        if best is None or key > best[:2]:
            best = (q, ci == prefer_chunk, ci, s, e, blocks)
    if best is None:
        return None
    return best[2], best[3], best[4], best[0], best[5]


# ---------------------------------------------------------------------------
# Inputs: the extraction JSON names its own paper, so both the text and the
# output filename are derived from the one argument the user passes.
# ---------------------------------------------------------------------------

def load_records(json_path: Path):
    """Read an extraction or merged JSON as [(record, reaction_dict), ...]."""
    data = json.loads(json_path.read_text())
    if isinstance(data, dict):                      # tolerate {"reactions": [...]}
        data = data.get("reactions") or data.get("results") or []
    out = []
    for rec in data:
        if not isinstance(rec, dict):
            continue
        # extraction/merged wrap the reaction in `annotation_result`; accept a
        # bare reaction dict too, so hand-filtered files still work.
        out.append((rec, rec.get("annotation_result") or rec))
    return out


def sources_of(records):
    """Distinct `source` values recorded in the JSON, in first-seen order."""
    seen = []
    for rec, _ in records:
        s = rec.get("source")
        if s and s not in seen:
            seen.append(s)
    return seen


MERGED_GLOB = "*_merged.json"


def merged_candidates():
    """Every results/*_merged.json, newest first."""
    return sorted(RESULTS_DIR.glob(MERGED_GLOB),
                  key=lambda p: p.stat().st_mtime, reverse=True)


def describe_candidates(cands, limit=None):
    """One line per merged file: age, reaction count and the paper it annotates."""
    import datetime
    for p in cands[:limit]:
        when = datetime.datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
        try:
            recs = load_records(p)
            what = f"{len(recs):3d} reaction(s)  {', '.join(sources_of(recs)) or '(no source)'}"
        except (json.JSONDecodeError, OSError) as e:
            what = f"unreadable ({type(e).__name__})"
        print(f"    {when}  {p.name:55s} {what}")


def latest_merged(source=None):
    """The newest merged JSON, optionally restricted to one paper.

    `source` matches the value the pipeline recorded in the JSON, so asking for
    PINK1.pdf will not silently hand back the PMC4003245 run of the same paper —
    those are different text routes and produce different reaction sets.
    """
    cands = merged_candidates()
    if not cands:
        raise SystemExit(
            f"no {MERGED_GLOB} in {RESULTS_DIR} — run the merge step first, or pass "
            f"an explicit path with --allow-unmerged to annotate a raw extraction")
    if not source:
        return cands[0]
    for p in cands:
        try:
            if source in sources_of(load_records(p)):
                return p
        except (json.JSONDecodeError, OSError):
            continue
    print(f"[err] no merged JSON records source {source!r}. Available:")
    describe_candidates(cands)
    raise SystemExit(1)


def get_results_text(source: str, use_cache=True, refresh=False, verbose=False):
    """The Results section for `source`, exactly as run_extraction.py loads it."""
    cache = CACHE_DIR / f"{_slug(source)}_results.txt"
    if use_cache and not refresh and cache.is_file():
        if verbose:
            print(f"[text] cache hit: {cache}")
        return cache.read_text(), f"cache:{cache.name}"

    # Imported lazily: this pulls in the Anthropic client (and, for a PDF, LLM
    # section detection), which the --text path has no business requiring.
    sys.path.insert(0, str(ROOT))
    sys.path.insert(0, str(ROOT / "reactome_llm"))
    from dotenv import load_dotenv
    load_dotenv(ROOT / ".env", override=True)
    import anthropic
    import PubMedFetcher as fetcher

    client = anthropic.Anthropic(api_key=os.getenv("ANTHROPIC_API_KEY"))
    _source_id, text, how = fetcher.load_source(
        source, client=client, model="claude-sonnet-5")
    if use_cache:
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        cache.write_text(text)
        if verbose:
            print(f"[text] cached to {cache}")
    return text, how


def chunk_text(text):
    """Split with run_extraction's splitter OBJECT, so these are the extraction's chunks.

    A local RecursiveCharacterTextSplitter with copied parameters would drift the
    moment run_extraction.py is tuned, and the docx would then annotate chunks the
    extraction never saw. run_extraction's CLI is behind an __main__ guard, so
    importing it only builds its client and graph.
    """
    sys.path.insert(0, str(ROOT))
    import run_extraction as pipeline
    return pipeline.splitter.split_text(text)


def _slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", str(s)).strip("_").lower()


def default_output(json_path: Path) -> Path:
    """results/<json stem>_annotated.docx — derived, never a hardcoded gene name."""
    return RESULTS_DIR / f"{json_path.stem}_annotated.docx"


def parse_args(argv=None):
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("merged", nargs="?",
                    help="results/<stem>_merged.json. Default: the newest one "
                         "(restricted to --source, if given)")
    ap.add_argument("--source",
                    help="paper to annotate (PDF name in data/papers/, PMID, or PMCID). "
                         "With no positional argument, also selects WHICH merged JSON "
                         "is used. Default: the `source` recorded in the JSON")
    ap.add_argument("--list", action="store_true",
                    help="list the merged JSONs, newest first, and exit")
    ap.add_argument("--allow-unmerged", action="store_true",
                    help="annotate a raw *_extraction.json anyway (refused by default: "
                         "it still contains the per-chunk duplicates the merge collapses)")
    ap.add_argument("--text",
                    help="use this plain-text file as the Results section instead of "
                         "fetching the paper (no PubMed request, no PDF section detection)")
    ap.add_argument("--out", help=f"output .docx (default: {default_output(Path('X.json')).name} "
                                  "pattern under results/)")
    ap.add_argument("--refresh", action="store_true",
                    help="re-fetch the Results text even if it is cached")
    ap.add_argument("--no-cache", action="store_true",
                    help="neither read nor write the Results-text cache")
    ap.add_argument("-v", "--verbose", action="store_true",
                    help="print a per-reaction match report (chunk, offset, quality)")
    return ap.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)

    if args.list:
        cands = merged_candidates()
        print(f"[merged] {len(cands)} file(s) in {RESULTS_DIR}, newest first:")
        describe_candidates(cands)
        return

    if args.merged:
        json_path = Path(args.merged)
        if not json_path.is_file():                 # also accept a bare stem/name
            for cand in (RESULTS_DIR / args.merged,
                         RESULTS_DIR / f"{args.merged}.json",
                         RESULTS_DIR / f"{args.merged}_merged.json"):
                if cand.is_file():
                    json_path = cand
                    break
            else:
                raise SystemExit(f"no such merged JSON: {args.merged}")
        chosen = "given on the command line"
    else:
        json_path = latest_merged(args.source)
        chosen = "newest merged JSON" + (f" for {args.source}" if args.source else "")

    # Refused rather than warned: a report built from an extraction shows the same
    # reaction once per chunk it was mentioned in, and nothing in the docx says so.
    if not json_path.name.endswith("_merged.json") and not args.allow_unmerged:
        print(f"[err] {json_path.name} is not a merged JSON. Merged files available:")
        describe_candidates(merged_candidates())
        raise SystemExit(
            "pass one of the above, run with no argument to take the newest, or "
            "--allow-unmerged to annotate this file anyway")

    records = load_records(json_path)
    if not records:
        raise SystemExit(f"{json_path} contains no reactions")

    json_sources = sources_of(records)
    if args.text:
        source = args.source or (json_sources[0] if json_sources else str(args.text))
        results_text = Path(args.text).read_text()
        how = f"file:{args.text}"
    else:
        source = args.source or (json_sources[0] if json_sources else None)
        if not source:
            raise SystemExit(
                f"{json_path} records no `source` — pass --source <PDF|PMID|PMCID> "
                f"or --text <file>")
        results_text, how = get_results_text(
            source, use_cache=not args.no_cache, refresh=args.refresh,
            verbose=args.verbose)

    if len(json_sources) > 1:
        print(f"[warn] JSON mixes {len(json_sources)} sources {json_sources}; "
              f"annotating against {source!r} only — reactions from the others "
              f"will report as unmatched")
    elif json_sources and args.source and json_sources[0] != args.source:
        print(f"[warn] --source {args.source!r} does not match the JSON's "
              f"source {json_sources[0]!r}")

    chunks = chunk_text(results_text)
    out_path = Path(args.out) if args.out else default_output(json_path)

    # Echoed up front: the whole point is that the reader can see which JSON,
    # which paper and which output file this run actually used.
    print(f"[in ] merged json: {json_path}   ({chosen})")
    print(f"[in ] paper      : {source}   (Results via {how})")
    print(f"[in ] text       : {len(results_text.split()):,} words -> {len(chunks)} chunks")
    print(f"[out] docx       : {out_path}")

    norm_chunks = [normalize(c) for c in chunks]

    # Per-chunk accumulators, keyed by chunk index.
    chunk_highlights = {i: [] for i in range(len(chunks))}   # [(start, end)]
    # one anchor per EVIDENCE SENTENCE, not per reaction, so every highlight
    # carries its own comment and a wrong-sentence match is visible on the spot
    chunk_anchors = {i: [] for i in range(len(chunks))}      # [(start, end, rxn, n, tot)]

    total_rxn = 0
    matched_rxn = 0
    total_comments = 0
    unmatched = []

    for rec, rxn in records:
        total_rxn += 1
        prefer = rec.get("chunk_index")
        ev_matches = []  # (chunk_idx, start, end, quality, blocks)
        seen_spans = set()
        for ev in rxn.get("evidence") or []:
            hit = locate_evidence(ev, norm_chunks, prefer_chunk=prefer)
            if hit is None:
                continue
            # two evidence strings can resolve to the same sentence; one comment
            # per span is enough, a duplicate would just crowd the margin
            if hit[:3] in seen_spans:
                continue
            seen_spans.add(hit[:3])
            ev_matches.append(hit)

        if ev_matches:
            matched_rxn += 1
            # numbered in reading order, so "evidence 2/3" means the second
            # sentence as it appears in the paper
            ev_matches.sort(key=lambda t: (t[0], t[1]))
            for n, (ci, s, e, _q, blocks) in enumerate(ev_matches, 1):
                # highlight only the aligned pieces — anything sitting between
                # them belongs to no evidence sentence and stays unhighlighted
                chunk_highlights[ci].extend(blocks)
                chunk_anchors[ci].append((s, e, rxn, n, len(ev_matches)))
                total_comments += 1
            if args.verbose:
                q = ", ".join(f"c{c}@{s} q={qq:.2f}" for c, s, _e, qq, _b in ev_matches)
                print(f"  MATCH   (json chunk {prefer}) {rxn.get('name')}  "
                      f"[{len(ev_matches)} comment(s): {q}]")
        else:
            unmatched.append((prefer, rxn.get("name")))
            if args.verbose:
                for ev in rxn.get("evidence") or []:
                    print(f"  NOMATCH (json chunk {prefer}) {rxn.get('name')}"
                          f"\n            evidence: {normalize(ev)[:120]}...")

    doc = Document()

    # ---- Title / preamble ----
    label = Path(str(source)).stem or str(source)
    doc.add_heading(f"{label} — RESULTS reactions annotated onto pipeline chunks", level=0)
    intro = doc.add_paragraph()
    intro.add_run(
        "Every evidence sentence an extracted reaction was pulled from is "
    )
    hl = intro.add_run("highlighted")
    hl.font.highlight_color = WD_COLOR_INDEX.YELLOW
    intro.add_run(
        " below and carries its OWN comment, so a reaction cited by three sentences "
        "appears three times, labelled 'evidence 1/3' and so on. Only the evidence "
        "sentences are highlighted — unhighlighted text produced no reaction. Source: "
    )
    intro.add_run(f"{json_path.name} / {source}").italic = True

    for ci, _chunk_text in enumerate(chunks):
        hay = norm_chunks[ci]
        highlight_spans = chunk_highlights[ci]
        anchors = chunk_anchors[ci]

        n_rxn = len({id(rxn) for _s, _e, rxn, _n, _t in anchors})
        doc.add_heading(f"Chunk {ci}   ({len(anchors)} comment(s), "
                        f"{n_rxn} reaction(s))", level=1)

        # Build the set of cut points to segment the paragraph into runs.
        cuts = {0, len(hay)}
        for s, e in highlight_spans:
            cuts.add(s)
            cuts.add(e)
        for s, e, _rxn, _n, _t in anchors:
            cuts.add(s)
            cuts.add(e)
        cuts = sorted(c for c in cuts if 0 <= c <= len(hay))

        para = doc.add_paragraph()
        seg_runs = []  # list of (seg_start, seg_end, run)
        for a, b in zip(cuts, cuts[1:]):
            if b <= a:
                continue
            run = para.add_run(hay[a:b])
            if any(s <= a and b <= e for s, e in highlight_spans):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            seg_runs.append((a, b, run))

        # One comment per evidence sentence, anchored to the runs covering that
        # sentence only — so clicking a highlight names the reaction it produced.
        for s, e, rxn, n, tot in anchors:
            runs = [run for (a, b, run) in seg_runs if a >= s and b <= e]
            if not runs:
                runs = [seg_runs[0][2]] if seg_runs else None
            if not runs:
                continue
            attach_comment(doc, runs, build_comment(rxn, ev_num=n, ev_total=tot))

    # ---- Summary section ----
    doc.add_heading("Annotation summary", level=1)
    p = doc.add_paragraph()
    p.add_run(f"Merged JSON: {json_path.name}\n")
    p.add_run(f"Paper: {source}  (Results via {how})\n")
    p.add_run(f"Chunks: {len(chunks)}\n")
    p.add_run(f"Reactions total: {total_rxn}\n")
    p.add_run(f"Anchored to at least one evidence sentence: {matched_rxn}\n")
    p.add_run(f"Comments placed (one per evidence sentence): {total_comments}\n")
    p.add_run(f"Reactions with no evidence found in the Results text: {len(unmatched)}")
    if unmatched:
        for ci, name in unmatched:
            doc.add_paragraph(f"(json chunk {ci}) {name}", style="List Bullet")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)

    print(f"Wrote {out_path}")
    print(f"Reactions: {total_rxn}  matched: {matched_rxn}  unmatched: {len(unmatched)}"
          f"  comments: {total_comments}")
    if not args.verbose:
        for ci, name in unmatched:
            print(f"  UNMATCHED (json chunk {ci}): {name}")


if __name__ == "__main__":
    main()
